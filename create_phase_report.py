from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

OUT = 'VIPCAR_Rapport_Phase_Projet_FR.docx'
GREEN, GOLD, INK, MUTED = '133A31', 'B48A46', '161814', '6D7068'
PALE, LIGHT = 'F7F5EF', 'F1F1EB'

def set_run(run, size=11, color=INK, bold=False, italic=False, font='Calibri'):
    run.font.name = font
    run._element.rPr.rFonts.set(qn('w:ascii'), font)
    run._element.rPr.rFonts.set(qn('w:hAnsi'), font)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold, run.italic = bold, italic

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd')) or OxmlElement('w:shd')
    if shd.getparent() is None: tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def cell_edges(cell, color='DCDCD4'):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders') or OxmlElement('w:tcBorders')
    if borders.getparent() is None: tcPr.append(borders)
    for edge in ('top','left','bottom','right'):
        el = borders.find(qn('w:'+edge)) or OxmlElement('w:'+edge)
        if el.getparent() is None: borders.append(el)
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'6'); el.set(qn('w:space'),'0'); el.set(qn('w:color'),color)

def cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = tcPr.first_child_found_in('w:tcMar') or OxmlElement('w:tcMar')
    if mar.getparent() is None: tcPr.append(mar)
    for name, value in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        el = mar.find(qn('w:'+name)) or OxmlElement('w:'+name)
        if el.getparent() is None: mar.append(el)
        el.set(qn('w:w'),str(value)); el.set(qn('w:type'),'dxa')

def table_width(table, widths):
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblW = tblPr.first_child_found_in('w:tblW') or OxmlElement('w:tblW')
    if tblW.getparent() is None: tblPr.append(tblW)
    tblW.set(qn('w:w'), str(sum(widths))); tblW.set(qn('w:type'),'dxa')
    grid = table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for width in widths:
        col = OxmlElement('w:gridCol'); col.set(qn('w:w'),str(width)); grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.first_child_found_in('w:tcW') or OxmlElement('w:tcW')
            if tcW.getparent() is None: tcPr.append(tcW)
            tcW.set(qn('w:w'),str(width)); tcW.set(qn('w:type'),'dxa')

def add_page_number(paragraph):
    run = paragraph.add_run()
    for kind, text in [('w:fldChar','begin'),('w:instrText',' PAGE '),('w:fldChar','end')]:
        node = OxmlElement(kind)
        if kind == 'w:fldChar': node.set(qn('w:fldCharType'), text)
        else: node.set(qn('xml:space'),'preserve'); node.text = text
        run._r.append(node)
    set_run(run, 9, MUTED)

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
section.header_distance = section.footer_distance = Inches(0.492)

normal = doc.styles['Normal']
normal.font.name='Calibri'; normal._element.rPr.rFonts.set(qn('w:ascii'),'Calibri'); normal._element.rPr.rFonts.set(qn('w:hAnsi'),'Calibri')
normal.font.size=Pt(11); normal.font.color.rgb=RGBColor.from_string(INK)
normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.10
for name,size,color,before,after in [('Heading 1',16,GREEN,16,8),('Heading 2',13,GREEN,12,6),('Heading 3',12,INK,8,4)]:
    st=doc.styles[name]; st.font.name='Calibri'; st._element.rPr.rFonts.set(qn('w:ascii'),'Calibri'); st._element.rPr.rFonts.set(qn('w:hAnsi'),'Calibri')
    st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color)
    st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after); st.paragraph_format.keep_with_next=True
if 'Kicker' not in [s.name for s in doc.styles]: doc.styles.add_style('Kicker', WD_STYLE_TYPE.PARAGRAPH)
doc.styles['Kicker'].font.name='Calibri'; doc.styles['Kicker'].font.size=Pt(9); doc.styles['Kicker'].font.bold=True; doc.styles['Kicker'].font.color.rgb=RGBColor.from_string(GOLD)

header=section.header.paragraphs[0]
r=header.add_run('VIPCAR TUNISIA  /  RAPPORT DE PHASE'); set_run(r,8.5,MUTED,True)
footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('Document client  •  Août 2026  •  Page '); set_run(r,9,MUTED); add_page_number(footer)

def paragraph(text='', style=None, align=WD_ALIGN_PARAGRAPH.LEFT, before=None, after=None):
    p=doc.add_paragraph(style=style); p.alignment=align
    if text:
        r=p.add_run(text); set_run(r)
    if before is not None: p.paragraph_format.space_before=Pt(before)
    if after is not None: p.paragraph_format.space_after=Pt(after)
    return p

def bullet(text):
    p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(4); p.paragraph_format.line_spacing=1.10
    r=p.add_run(text); set_run(r); return p

def callout(label,text):
    t=doc.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; table_width(t,[9360])
    c=t.cell(0,0); shade(c,LIGHT); cell_edges(c,'D7D8D1'); cell_margins(c,150,180,150,180)
    p=c.paragraphs[0]; r=p.add_run(label.upper()); set_run(r,9,GOLD,True)
    p=c.add_paragraph(); p.paragraph_format.space_after=Pt(0); r=p.add_run(text); set_run(r,10.5)
    return t

def status_table(rows):
    t=doc.add_table(rows=1,cols=3); t.alignment=WD_TABLE_ALIGNMENT.CENTER; table_width(t,[1800,1800,5760])
    for cell,text in zip(t.rows[0].cells,['Domaine','Statut','Résultat / observation']):
        shade(cell,GREEN); cell_edges(cell,GREEN); cell_margins(cell)
        r=cell.paragraphs[0].add_run(text); set_run(r,9,'FFFFFF',True)
    for domain,status,detail in rows:
        cells=t.add_row().cells
        for cell in cells: cell_edges(cell); cell_margins(cell)
        shade(cells[1], 'E5F0EA' if status=='Terminé' else 'FFF3D6')
        for cell,text in zip(cells,[domain,status,detail]):
            r=cell.paragraphs[0].add_run(text); set_run(r,9.5,GREEN if text==status and status=='Terminé' else INK,text==status)
    return t

p=paragraph('RAPPORT DE PHASE','Kicker',after=8)
p=paragraph('Projet VIPCAR Tunisia',after=3); p.paragraph_format.space_before=Pt(20); set_run(p.runs[0],28,GREEN,True)
p=paragraph('Bilan des travaux réalisés et niveau d’avancement',after=20); set_run(p.runs[0],15,MUTED)
meta=doc.add_table(rows=4,cols=2); meta.alignment=WD_TABLE_ALIGNMENT.LEFT; table_width(meta,[1800,7560])
for row,(label,value) in zip(meta.rows,[('Client','VIPCAR Tunisia'),('Périmètre','Site web, UX/UI, flotte, SEO et responsive'),('Phase','Construction, intégration et optimisation'),('Date','23 août 2026')]):
    for cell in row.cells: cell_edges(cell); cell_margins(cell,85,120,85,120)
    shade(row.cells[0],PALE); r=row.cells[0].paragraphs[0].add_run(label); set_run(r,9,GREEN,True)
    r=row.cells[1].paragraphs[0].add_run(value); set_run(r,10.5)
paragraph('',after=12)
callout('Synthèse exécutive','La phase a permis de transformer le projet VIPCAR en une plateforme bilingue de mobilité en Tunisie, avec une flotte enrichie, des parcours de réservation, des pages locales orientées SEO et une interface responsive revue sur desktop, tablette et mobile.')

doc.add_heading('1. Objectif de la phase',level=1)
paragraph('L’objectif était de construire une base digitale premium pour VIPCAR, capable de présenter clairement les services de location de voiture, de transferts aéroport et de chauffeur privé, tout en facilitant la prise de contact et la demande de devis via WhatsApp.')
for item in ['Structurer une expérience cohérente en anglais et en français.','Présenter la flotte et les catégories de véhicules de manière claire.','Développer des pages services, destinations et contenus éditoriaux utiles.','Poser une base SEO technique et éditoriale exploitable pour la suite.']: bullet(item)

doc.add_heading('2. Travaux réalisés',level=1)
doc.add_heading('2.1 Architecture et navigation',level=2)
paragraph('Une architecture complète a été mise en place autour des principaux besoins du parcours client : accueil, flotte, services, réservation, destinations, blog, corporate, contact, FAQ et pages légales.')
for item in ['Navigation bilingue avec changement de langue conservant la page consultée.','Routes dédiées pour les véhicules, les villes, les aéroports et les services avec chauffeur.','Header, menu mobile, CTA de réservation et liens WhatsApp intégrés à l’ensemble du site.']: bullet(item)

doc.add_heading('2.2 Flotte et fiches véhicules',level=2)
paragraph('La flotte a été enrichie et structurée autour de 18 véhicules, répartis entre modèles économiques, compacts, berlines, SUV, premium, pick-up et véhicules de groupe.')
status_table([('Flotte','Terminé','18 véhicules intégrés avec catégorie, prix indicatif, places, bagages et transmission.'),('Filtres','Terminé','Filtrage par catégorie avec interface adaptée au mobile.'),('Fiches détail','Terminé','Page dédiée pour chaque véhicule avec demande de devis contextualisée.')])

doc.add_heading('2.3 Services et conversion',level=2)
paragraph('Les trois offres principales ont été développées : location de voiture, transferts aéroport et chauffeur privé. Chaque service dispose d’un contenu explicatif, d’un formulaire de devis et de liens vers les pages associées.')
for item in ['Formulaire de demande de devis connecté au parcours WhatsApp.','Sélection du type de service, lieu de départ, date et informations complémentaires.','CTA répétés aux endroits stratégiques sans multiplier les étapes de réservation.']: bullet(item)

doc.add_heading('2.4 Expérience utilisateur et design',level=2)
paragraph('Plusieurs passes d’amélioration UI ont été réalisées pour aligner l’interface avec une identité premium : typographie éditoriale, palette sombre et dorée, cartes, états actifs, espacements et hiérarchie visuelle.')
for item in ['Refonte du sélecteur de langue avec menu personnalisé EN / FR.','Amélioration des filtres de flotte et des cartes véhicules.','Refonte des blocs de statistiques, de la section About et des blocs de demande de devis.','Correction des espacements et fonds de la page Destinations.','Ajustements du menu hamburger et des CTA mobiles.']: bullet(item)

doc.add_heading('2.5 Responsive et contrôle qualité',level=2)
paragraph('Les principaux parcours ont été vérifiés sur des largeurs desktop, tablette et mobile. Les problèmes d’overflow horizontal, de grilles trop denses, de formulaires inline et de CTA masqués ont été traités.')
status_table([('Mobile','Terminé','Header, menu, sélecteur de langue, formulaires et CTA adaptés.'),('Tablette','Terminé','Grilles destinations, journal, flotte et sections de contenu rééquilibrées.'),('Vérification','Terminé','Routes clés contrôlées sans débordement horizontal.')])

doc.add_heading('3. SEO et visibilité',level=1)
paragraph('Une fondation SEO dynamique a été intégrée directement dans l’application. Les métadonnées sont générées selon la langue et la route consultée.')
for item in ['Titres et descriptions dynamiques pour les pages principales, services, destinations, articles et véhicules.','URL canoniques, Open Graph, Twitter Card et liens hreflang anglais/français.','Données structurées LocalBusiness / AutoRental, Service, FAQ, Breadcrumb et Vehicle / Product.','Sitemap XML et robots.txt présents.','Pages éditoriales sur la location, l’aéroport, la conduite, les documents et la location longue durée.']: bullet(item)
callout('Point à finaliser','Le fichier de mots-clés annoncé n’était pas présent dans l’espace de travail au moment de l’audit. La stratégie de mots-clés, le mapping page par page et les recommandations de contenu restent à compléter dès réception du fichier.')

doc.add_heading('4. Résultats atteints',level=1)
status_table([('Plateforme bilingue','Terminé','Parcours EN / FR avec contenus, navigation et métadonnées localisés.'),('Catalogue véhicules','Terminé','18 véhicules visibles et filtrables.'),('Pages de conversion','Terminé','Services, booking, fiches véhicules et CTA WhatsApp.'),('Pages locales','Terminé','Destinations, location, transferts et chauffeur par zone.'),('Base SEO','Terminé','Métadonnées, schema, sitemap et robots intégrés.'),('Mots-clés','À recevoir','Analyse détaillée en attente du fichier client.')])

doc.add_heading('5. Points de vigilance',level=1)
for item in ['Le sitemap et le script de pré-rendu doivent être synchronisés avec les cinq véhicules ajoutés récemment.','Le pré-rendu complet dépend du lancement de Chromium via Playwright ; le build client Vite est valide, mais le pré-rendu local a rencontré une restriction de permission dans l’environnement de travail.','Les images sont actuellement servies depuis le domaine VIPCAR ; une stratégie d’assets optimisés et de fallback peut encore améliorer la robustesse et les performances.','La validation finale des données structurées et des balises SEO doit être réalisée sur l’environnement de production.']: bullet(item)

doc.add_heading('6. Prochaine phase recommandée',level=1)
for item in ['Intégrer le fichier de mots-clés et réaliser le mapping intention → page.','Mettre à jour sitemap.xml et le script de pré-rendu pour les 18 véhicules.','Finaliser les titres, H1, descriptions et liens internes selon les priorités SEO.','Optimiser les images, les performances et les Core Web Vitals.','Effectuer une recette client finale sur les parcours réservation, WhatsApp, langue et mobile.','Préparer le déploiement et le suivi analytics de production.']: bullet(item)
callout('Conclusion','La phase a atteint son objectif de construction et de mise en qualité de la plateforme VIPCAR. Le site dispose maintenant d’une base fonctionnelle, bilingue, responsive et orientée conversion. La prochaine étape consiste à aligner cette base avec les mots-clés prioritaires et à finaliser la préparation production.')

doc.core_properties.title='Rapport de phase - Projet VIPCAR Tunisia'
doc.core_properties.subject='Bilan des travaux réalisés et niveau d’avancement'
doc.core_properties.author='VIPCAR Tunisia'
doc.save(OUT)
print(OUT)

